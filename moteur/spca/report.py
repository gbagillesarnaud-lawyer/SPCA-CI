"""Module 9 — Report Generator (rapport SOLEX A–H, Markdown et Word)."""
from __future__ import annotations

from .common import MSG_INFO, MSG_NQ, fmt

STATUS_FR = {"CONFORME": "CONFORME", "NON_CONFORME": "NON CONFORME", "NON_APPLICABLE": "NON APPLICABLE"}


def _st(c):
    s = STATUS_FR.get(c["status"], c["status"])
    return s + (f" — {c['severity']}" if c.get("severity") else "")


def _val(v, missing=MSG_NQ):
    return fmt(v) if v is not None else missing


def sections(r: dict) -> list[tuple[str, list]]:
    """Structure commune Markdown / Word : [(titre, [blocs])], bloc = ('p', texte) | ('t', entêtes, lignes)."""
    e, emp, ctx, rc = r["employee"], r["employer"], r["applicable_context"], r["payroll_reconstruction"]
    rs = r["risk_summary"]
    ctrl = [c for c in r["controls"] if c["status"] != "NON_APPLICABLE"]
    S = []
    S.append(("A. Identification", [("t", ["Élément", "Valeur"], [
        ["Entreprise", emp.get("legal_name") or MSG_INFO], ["Secteur", emp.get("sector") or MSG_INFO],
        ["Salarié (ID anonymisé)", e.get("employee_id")], ["Fonction", e.get("position") or MSG_INFO],
        ["Classification / catégorie", f"{e.get('classification') or '—'} / {e.get('category') or '—'}"],
        ["Convention collective", ctx.get("collective_agreement") or MSG_INFO],
        ["Période", r["payroll_period"] or MSG_INFO], ["Date d'embauche", e.get("hire_date") or MSG_INFO],
        ["Régime de durée du travail", ctx.get("working_time_regime") or MSG_INFO],
        ["Version agent / référentiel", f"{r['audit_log']['agent_version']} / {r['audit_log']['ruleset_version']}"],
    ])]))
    b = [("p", f"**Statut : {r['status']}** — {r['report_status']}"),
         ("p", f"P1 : {rs['P1']} · P2 : {rs['P2']} · P3 : {rs['P3']} · P4 : {rs['P4']} · "
               f"Niveau de confiance : **{r['confidence']}** ({r['confidence_score']})")]
    cs = r.get("compliance_score") or {}
    if cs.get("score") is not None:
        b.append(("p", f"Score de conformité SOLEX : {cs['score']} % — {cs['qualification']} "
                       f"({cs['controls_counted']} contrôles conclusifs ; indicateur méthodologique interne, distinct de la qualification juridique)"))
    b += [("p", f"⛔ {s}") for s in r["stops"]] + [("p", f"• {n}") for n in r["notes"]]
    nb_nc = sum(1 for c in ctrl if c["status"] not in ("CONFORME", "NON_CONFORME"))
    b.append(("p", f"{len(ctrl)} contrôles exécutés, dont {sum(1 for c in ctrl if c['status'] == 'NON_CONFORME')} non-conformités "
                   f"et {nb_nc} points non contrôlés ou non déterminés."))
    S.append(("B. Conclusion exécutive", b))
    S.append(("C. Tableau des contrôles", [("t", ["#", "Élément contrôlé", "Constaté", "Attendu", "Écart", "Référence", "Statut", "Correction"],
              [[f"{c['control_id']}", c["control_name"], fmt(c["observed_value"]), fmt(c["expected_value"]), fmt(c["difference"]),
                c["legal_reference"], _st(c), c["recommendation"]] for c in ctrl]),
              ("p", "**Constats, réserves et traces de calcul**")] +
             [("p", f"**{c['control_id']} {c['control_name']}** — {c['finding'] or _st(c)}"
                    + (f" · Réserves : {' ; '.join(c['reserves'])}" if c["reserves"] else "")
                    + "".join(f" · [{t.get('step')}] {t.get('formula')} ⇒ {t.get('result')}"
                              + (f" (entrées : {t['inputs']})" if t.get("inputs") else "")
                              + (f" (détail : {' ; '.join(t['steps'])})" if t.get("steps") else "") for t in c["trace"]))
              for c in ctrl]))
    S.append(("D. Reconstitution de la paie", [("t", ["Poste", "Montant"], [
        ["Brut bulletin", _val(rc.get("gross_reported"), MSG_INFO)], ["Brut recalculé", _val(rc.get("gross_recalculated"))],
        ["Assiette CNPS (bulletin)", _val(r.get("inputs") and next((f["value"] for f in r["inputs"]["fields"] if f["field"] == "payroll.social_base"), None), MSG_INFO)],
        ["CNPS salariale théorique", _val(rc.get("cnps_employee_theoretical"))], ["CMU théorique", _val(rc.get("cmu_employee_theoretical"))],
        ["ITS théorique", _val(rc.get("its_theoretical"))], ["Autres retenues", _val(rc.get("other_deductions"))],
        ["Net théorique", _val(rc.get("net_theoretical"), rc.get("net_theoretical_status", MSG_NQ))],
        ["Net bulletin", _val(rc.get("net_reported"), MSG_INFO)], ["Écart net", _val(rc.get("net_difference"))],
    ])]))
    f = r["financial_summary"]
    rows = [[k.replace("_", " "), fmt(v)] for k, v in f.items() if k not in ("non_quantifiables", "note")]
    S.append(("E. Synthèse financière", [("t", ["Poste", "Montant (mois audité)"], rows or [["Aucun écart quantifié", "0"]]),
              ("p", "Montants non quantifiables : " + (", ".join(f["non_quantifiables"]) or "aucun")), ("p", f["note"])]))
    S.append(("F. Plan correctif", [("t", ["Anomalie", "Correction", "Responsable", "Priorité", "Échéance", "Régularisation rétroactive"],
              [[a["anomaly"], a["correction"], a["owner"], a["priority"], a["due_date"], a["retroactive_regularisation"]]
               for a in r["corrective_actions"]] or [["Aucune anomalie sur les points contrôlés", "—", "—", "—", "—", "—"]])]))
    g = [("p", f"**{r['escalation_formula']}**")] if r["escalation_formula"] else [("p", "Aucune escalade automatique.")]
    if r["human_escalations"]:
        g.append(("t", ["Code", "Motif", "Priorité", "Expert", "Actions bloquées"],
                  [[x["escalation_type"], x["reason"], x["priority"], x["recommended_expert"], ", ".join(x["blocked_actions"]) or "—"]
                   for x in r["human_escalations"]]))
    S.append(("G. Escalades", g))
    src = [("p", f"• {s.get('rule_id') or s.get('parameter')} {('v' + str(s['version'])) if s.get('version') else ''} — "
                 f"{s.get('legal_source')}{(', art. ' + str(s['article'])) if s.get('article') else ''} — effet {s.get('effective_from')}"
                 + (f" — {s['source_url']}" if s.get("source_url") else "")) for s in r["sources"]] or [("p", "Aucune source appliquée.")]
    S.append(("H. Sources", src))
    return S


def to_markdown(r: dict) -> str:
    L = [f"# Rapport d'audit de paie — {r['audit_id']}", ""]
    for title, blocks in sections(r):
        L += [f"## {title}", ""]
        for bl in blocks:
            if bl[0] == "p":
                L += [bl[1], ""]
            else:
                L.append("| " + " | ".join(bl[1]) + " |")
                L.append("|" + "---|" * len(bl[1]))
                L += ["| " + " | ".join(str(x).replace("|", "/").replace("\n", " ") for x in row) + " |" for row in bl[2]]
                L.append("")
    L += ["---", "_Aide à la décision générée par SPCA-CI. L'absence d'anomalie détectée ne vaut pas certification. "
          "Validation par un expert SOLEX obligatoire avant toute utilisation ou transmission._"]
    return "\n".join(L)


def to_docx(r: dict, path):
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        return None
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(9.5)
    d.add_heading(f"Rapport d'audit de paie — {r['audit_id']}", 0)
    for title, blocks in sections(r):
        d.add_heading(title, 1)
        for bl in blocks:
            if bl[0] == "p":
                d.add_paragraph(bl[1].replace("**", ""))
            else:
                t = d.add_table(rows=1, cols=len(bl[1]))
                t.style = "Table Grid"
                for i, h in enumerate(bl[1]):
                    t.rows[0].cells[i].text = h
                    for run in t.rows[0].cells[i].paragraphs[0].runs:
                        run.bold = True
                for row in bl[2]:
                    cells = t.add_row().cells
                    for i, v in enumerate(row):
                        cells[i].text = str(v)
    d.add_paragraph("Aide à la décision générée par SPCA-CI. Validation par un expert SOLEX obligatoire.").italic = True
    d.save(str(path))
    return path
